"""Extract clean text from uploaded files, keeping page/heading structure."""
from __future__ import annotations

import csv
import io
import json
import re
from html.parser import HTMLParser
from pathlib import Path


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        if tag in ("p", "br", "div", "li", "h1", "h2", "h3", "h4", "tr"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _blocks_from_text(text: str, page: int = 1) -> list[dict]:
    """Split into paragraph-level blocks so the chunker can pack them."""
    text = _clean(text)
    if not text:
        return []
    parts = [p.strip() for p in re.split(r"\n\s*\n", text)]
    out: list[dict] = []
    for p in parts:
        if not p:
            continue
        # very long paragraphs get split at line breaks, then sentences
        if len(p) > 1800:
            for line in [x.strip() for x in p.split("\n") if x.strip()]:
                out.extend({"text": s.strip(), "page": page} for s in _split_sentences(line, 1200))
        else:
            out.append({"text": p, "page": page})
    return out


def _split_sentences(text: str, max_len: int) -> list[str]:
    sentences = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9(\"'])", text)
    pieces, buf = [], ""
    for s in sentences:
        while len(s) > max_len:  # pathological long sentence
            cut = s.rfind(" ", 0, max_len) or max_len
            pieces.append(s[:cut])
            s = s[cut:].lstrip()
        if len(buf) + len(s) + 1 <= max_len:
            buf = f"{buf} {s}".strip()
        else:
            if buf:
                pieces.append(buf)
            buf = s
    if buf:
        pieces.append(buf)
    return pieces


# --------------------------------------------------------------------- PDF
HEADING_LIKE = re.compile(r"^(?:\d+(?:\.\d+)*[.)]?\s+)?[A-Z][^.!?,;:]{1,58}$")


def _separate_pdf_headings(text: str) -> str:
    """PDF extractors keep hard line breaks; promote short unpunctuated lines
    that precede a capitalised line into standalone blocks."""
    lines = [ln.rstrip() for ln in text.split("\n")]
    out: list[str] = []
    for i, line in enumerate(lines):
        out.append(line)
        stripped = line.strip()
        if not stripped or stripped.count("|") >= 2:
            continue
        nxt = lines[i + 1].strip() if i + 1 < len(lines) else ""
        if (HEADING_LIKE.match(stripped)
                and nxt and (nxt[0].isupper() or nxt[0].isdigit())
                and not stripped.endswith((".", ":", ";" , ","))
                and len(stripped.split()) <= 8):
            out.append("")  # blank line → own paragraph → own heading unit
    return "\n".join(out)


def parse_pdf(path: Path) -> tuple[list[dict], int]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    blocks: list[dict] = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if not text.strip():
            continue
        blocks.extend(_blocks_from_text(_separate_pdf_headings(text), page=i))
    return blocks, len(reader.pages)


# -------------------------------------------------------------------- DOCX
def parse_docx(path: Path) -> tuple[list[dict], int]:
    import docx

    document = docx.Document(str(path))
    lines: list[str] = []
    for para in document.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name or "").lower() if para.style else ""
        if "heading" in style or "title" in style:
            lines.append(f"\n## {t}\n")
        else:
            lines.append(t)
    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            lines.append("\n".join(rows))
    text = _clean("\n".join(lines))
    blocks = _blocks_from_text(text)
    pages = max(1, round(len(text) / 3000))
    return blocks, pages


# --------------------------------------------------------------------- CSV
def parse_csv(path: Path) -> tuple[list[dict], int]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    sample = raw[:4000]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except Exception:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(raw), dialect))
    if not rows:
        return [], 1
    header = [h.strip() for h in rows[0]]
    blocks: list[dict] = []
    for r in rows[1:]:
        cells = [v.strip() for v in r]
        if not any(cells):
            continue
        label = cells[0] if cells else ""
        pairs = ", ".join(f"{header[i] if i < len(header) else f'col{i}'}: {v}"
                          for i, v in enumerate(cells[1:], start=1) if v)
        text = f"{label} — {pairs}" if label else pairs
        if text and not text.endswith("."):
            text += "."
        if text.strip():
            blocks.append({"text": text.strip(), "page": 1})
    return blocks or [{"text": raw[:5000], "page": 1}], 1


# -------------------------------------------------------------------- HTML
def parse_html(path: Path) -> tuple[list[dict], int]:
    raw = path.read_text(encoding="utf-8", errors="replace")
    parser = _HTMLText()
    parser.feed(raw)
    text = _clean("".join(parser.parts))
    return _blocks_from_text(text), max(1, round(len(text) / 3000))


# ------------------------------------------------------------- plain / JSON
def parse_plain(path: Path) -> tuple[list[dict], int]:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except Exception:
            continue
    else:
        text = raw.decode("utf-8", errors="replace")
    if path.suffix.lower() == ".json":
        try:
            text = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
        except Exception:
            pass
    return _blocks_from_text(text), max(1, round(len(text) / 3000))


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".csv": parse_csv,
    ".html": parse_html,
    ".htm": parse_html,
    ".txt": parse_plain,
    ".md": parse_plain,
    ".markdown": parse_plain,
    ".json": parse_plain,
    ".log": parse_plain,
}


def parse_file(path: Path) -> tuple[list[dict], int]:
    ext = path.suffix.lower()
    fn = PARSERS.get(ext)
    if not fn:
        raise ValueError(f"Unsupported file type: {ext}")
    blocks, pages = fn(path)
    blocks = [b for b in blocks if b.get("text", "").strip()]
    if not blocks:
        raise ValueError("No extractable text found (is it a scanned/image PDF?)")
    return blocks, pages
